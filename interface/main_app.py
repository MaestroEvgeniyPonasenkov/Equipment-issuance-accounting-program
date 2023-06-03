import json
import sys

sys.path.insert(0, 'C:/Users/Stepan/PycharmProjects/pythonProject1/Equipment-issuance-accounting-program/')

import qdarkstyle
from PyQt5 import QtWidgets, QtCore
from PyQt5.QtWidgets import QFileDialog, QMessageBox, QTableWidgetItem, QListWidgetItem, QDialog, QDialogButtonBox, \
    QComboBox
from entry_ui import Ui_Enter
from window_ui import Ui_MainWindow
from user_ui import Ui_Newuser
from hardware_ui import Ui_Newhard
from alternative_ui import Ui_Alternative
from request_ui import Ui_NewRequest
from email_ui import Ui_Email

import alternative
import helpers
from xlwt import Workbook
from docx import Document
import requests
import re
import dotenv

CODE = '1'
DB_ACCESS_TOKEN = "Basic NVJOWUJkTGR1VER4UUNjTThZWXJiNW5BOkg0ZFNjQXlHYlM4OUtnTGdaQnMydlBzaw=="
DB_URL = "https://helow19274.ru/aip/api"


class LoginWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.table_window = Ui_MainWindow
        self.ui = Ui_Enter()
        self.ui.setupUi(self)
        self.ui.btn_entry.clicked.connect(self.on_click)
        self.ui.line_password.returnPressed.connect(self.on_click)

    def on_click(self):
        """Проверка пароля"""
        password = self.ui.line_password.text()
        if len(password) > 0:
            if password == CODE:
                self.close()
                self.table_window = MainWindow()
                self.table_window.show()
            else:
                self.ui.label.setText('Wrong password!')
                self.ui.line_password.clear()


class HardDialog(QDialog):
    def __init__(self):
        super().__init__()
        self.ui = Ui_Newhard()
        self.ui.setupUi(self)
        self.setStyleSheet(open('hard_style.qss').read())
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)

    def check(self):
        name = self.ui.name_line.text()
        specs = self.ui.spec_line.text()
        count = self.ui.count_line.text()
        description = self.ui.description_line.text()
        typ = self.ui.type_line.text()
        hardware_data = {
            "name": name,
            "type": typ,
            "description": description,
            "image_link": "https://encrypted-tbn0.gstatic.com/images?q=tbn:ANd9GcStW9UvyhB2beq-tiyJMhzWdP98Rny7PzRaPA&usqp=CAU",
            "specifications": specs
        }

        hardware_response = helpers.post_request('hardware', hardware_data)
        if 'detail' in hardware_response:
            if 'type' in hardware_response['detail'][0]['loc']:
                QMessageBox.warning(
                    self,
                    "Ошибка",
                    "Ошибка в вводе типа платы",
                    QMessageBox.StandardButton.Ok
                )
            elif 'name' in hardware_response['detail'][0]['loc']:
                QMessageBox.warning(
                    self,
                    "Ошибка",
                    "Ошибка в вводе названия",
                    QMessageBox.StandardButton.Ok
                )
            elif 'description' in hardware_response['detail'][0]['loc']:
                QMessageBox.warning(
                    self,
                    "Ошибка",
                    "Ошибка в вводе описания. Если оно отсутствует введите -",
                    QMessageBox.StandardButton.Ok
                )
            elif 'specifications' in hardware_response['detail'][0]['loc']:
                QMessageBox.warning(
                    self,
                    "Ошибка",
                    "Ошибка в вводе спецификации",
                    QMessageBox.StandardButton.Ok
                )
            else:
                print("Еще какая-то ошибка")
                print(hardware_response)
        else:
            self.close()


class UserDialog(QDialog):
    def __init__(self):
        super().__init__()
        self.ui = Ui_Newuser()
        self.ui.setupUi(self)
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)

    def check(self):
        surname = self.ui.surname_line.text()
        name = self.ui.name_line.text()
        patr = self.ui.patr_line.text()
        group = self.ui.group_line.text()
        email = self.ui.email_line.text()
        phone = self.ui.phone_line.text()
        access = self.ui.access_box.currentText()
        user_data = {
            "active": True,
            "type": access,
            "first_name": name,
            "last_name": surname,
            "patronymic": patr,
            "image_link": "https://clck.ru/34TfSF",
            "email": email,
            "phone": phone,
            "card_id": "string",
            "card_key": "string",
            "comment": group
        }

        user_request_body = json.dumps(user_data, ensure_ascii=False)

        response = requests.post(f"{DB_URL}/user",
                                 headers={'Authorization': DB_ACCESS_TOKEN},
                                 data=user_request_body
                                 )
        user_response = response.json()
        if 'detail' in user_response:
            if 'email' in user_response['detail'][0]['type']:
                QMessageBox.warning(
                    self,
                    "Ошибка",
                    "Ошибка в вводе почты",
                    QMessageBox.StandardButton.Ok
                )
            else:
                print("Еще какая-то ошибка")
                print(user_response)
        else:
            self.close()


class AlternativeDialog(QDialog):
    def __init__(self):
        super().__init__()
        self.ui = Ui_Alternative()
        self.ui.setupUi(self)
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.make_json)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)

    def make_json(self):
        log_value = int(self.ui.log_spin.text())
        log_index = float(self.ui.log_spin2.text().replace(',', '.'))
        mem_value = int(self.ui.mem_spin.text())
        mem_index = float(self.ui.mem_spin2.text().replace(',', '.'))
        mult_value = int(self.ui.mult_spin.text())
        mult_index = float(self.ui.mult_spin2.text().replace(',', '.'))
        pll_value = int(self.ui.pll_spin.text())
        pll_index = float(self.ui.pll_spin2.text().replace(',', '.'))
        pin_value = int(self.ui.pin_spin.text())
        pin_index = float(self.ui.pin_spin2.text().replace(',', '.'))
        data = {
            "log_elems": [log_value, log_index],
            "memory": [mem_value, mem_index],
            "pll": [pll_value, pll_index],
            "multiplier": [mult_value, mult_index],
            "pins": [pin_value, pin_index]
        }
        with open('max_variance.json', 'w', encoding='utf-8') as file:
            json.dump(data, file, indent=4, ensure_ascii=False)


class RequestDialog(QDialog):
    def __init__(self):
        super().__init__()
        self.ui = Ui_NewRequest()
        self.ui.setupUi(self)
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)

    def check(self):
        name = self.ui.name_line.text()
        last_name = self.ui.last_name_line.text()
        email = self.ui.email_line.text()
        phone = self.ui.phone_line.text()
        cabinet = self.ui.cab_line.text()
        board = self.ui.hardware_box.currentText()
        count = self.ui.count_box.text()
        issue_date = self.ui.date1_edit.text()
        return_date = self.ui.date2_edit.text()
        comment = self.ui.comment_line.text()
        data = {
            "user": 1,
            "location": cabinet,
            "status": "new",
            "comment": comment,
            "taken_date": issue_date,
            "return_date": return_date,
            "issued_by": 0,
            "hardware": [
                {
                    "hardware": board,
                    "count": count
                }
            ]
        }
        response = helpers.post_request('request', data)
        print(response)


def is_valid_email(email):
    pattern = r'^[a-zA-Z0-9._%+-]+@yandex.ru'
    return bool(re.match(pattern, email))


class Email(QDialog):
    def __init__(self):
        super().__init__()
        self.table_window = Ui_MainWindow
        self.ui = Ui_Email()
        self.ui.setupUi(self)
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)

    def check(self):
        """Проверка почты"""
        email = self.ui.email_line.text()
        password = self.ui.passwod_line.text()
        if (is_valid_email(email)):
            dotenv.set_key('.env', "EMAIL_USERNAME", email)
            dotenv.set_key('.env', "EMAIL_PASSWORD", password)
            print('ok')
        else:
            return QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе почты",
                QMessageBox.StandardButton.Ok
            )


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.email_dialog = Ui_Email
        self.status_box = QComboBox()
        self.request_dialog = Ui_NewRequest
        self.alter_dialog = Ui_Alternative
        self.hard_dialog = Ui_Newhard
        self.user_dialog = Ui_Newuser
        self.ui = Ui_MainWindow()
        self.ui.setupUi(self)
        self.ui.action_excel.triggered.connect(self.save_excel_file)
        self.ui.action_doc.triggered.connect(self.save_doc_file)
        self.ui.action_theme.triggered.connect(self.switch_theme)
        self.ui.action_alternative.triggered.connect(self.set_alternatives)
        self.ui.action_email.triggered.connect(self.configure_email)

        self.ui.users_button.clicked.connect(self.user_table)
        self.ui.hardware_button.clicked.connect(self.hardware_table)
        self.ui.request_button.clicked.connect(self.request_table)
        self.ui.list_button.clicked.connect(self.check_availability)

        self.ui.add_user_button.clicked.connect(self.add_user)
        self.ui.add_hardware_button.clicked.connect(self.add_hardware)
        self.ui.add_request_button.clicked.connect(self.add_request)

        self.ui.search_button.clicked.connect(self.search_user)
        self.ui.del_button.clicked.connect(self.delete)
        self.ui.tableWidget.itemChanged.connect(self.save_edits)
        self.boxes = {}

    def check_availability(self):
        """
        Проверка наличия платы
        """
        current_item = self.ui.hardware_list.currentItem()
        if current_item is None:
            return QMessageBox.warning(self, 'Внимание', 'Пожалуйста выберете плату для проверки')
        hardware_name = current_item.text()
        hardwares = helpers.get_request('hardware')
        stock = helpers.get_request('stocks')
        print(stock)
        print(hardware_name)
        hw_id = None
        current_hardware = {}
        for hw in hardwares:
            if hw.get('name') == hardware_name:
                hw_id = hw.get('id')
                current_hardware = hw
        print(hw_id)
        if hw_id is None:
            print("Ошибка")
        count = 0
        for st in stock:
            st_id = st['hardware']
            if st_id == hw_id:
                count += st['count']
        if count == 0:
            print("Данных плат нет в наличии")
            alter = alternative.find_alternative_board(current_hardware, hardwares)
            print(alter)

            message = QMessageBox()
            message.setWindowTitle("Наличие платы")
            if alter != '':
                message.setText(f'Данной платы нет в наличии, но доступна альтернатива:\n{alter}')
            else:
                message.setText(f'Данной платы нет в наличии, а также нет доступных альтернатив')
            message.setIcon(QMessageBox.Icon.Information)
            message.setStandardButtons(QMessageBox.Ok)
            message.exec_()

        else:
            print("Доступно", count, "плат")

    def configure_email(self):
        self.email_dialog = Email()
        self.email_dialog.show()

    def add_request(self):
        self.request_dialog = RequestDialog()
        self.hardware_table()
        self.request_dialog.show()
        boards = [self.ui.hardware_list.item(x).text() for x in range(self.ui.hardware_list.count())]
        self.request_dialog.ui.hardware_box.addItems(boards)

    def set_alternatives(self):
        self.alter_dialog = AlternativeDialog()
        self.alter_dialog.show()

    # Нужно изменить
    def save_edits(self):
        if self.ui.tableWidget.currentItem() is not None:
            item = self.ui.tableWidget.currentItem().text()
            current_row = self.ui.tableWidget.currentRow()
            current_col = self.ui.tableWidget.currentColumn()
            if self.ui.tableWidget.columnCount() == 8:
                user_dict = {"Имя": 'first_name', "Фамилия": 'last_name', "Отчество": 'patronymic',
                             "Группа": 'comment', "Уровень доступа": 'type', "Телефон": 'phone', "Почта": 'email'}

                col_name = user_dict[self.ui.tableWidget.horizontalHeaderItem(current_col).text()]

                update_id = int(self.ui.tableWidget.item(current_row, 0).text())
                user_response = helpers.patch_request('user', col_name, item, update_id)
                print(user_response)
                if user_response['detail'] == "OK":
                    QMessageBox.information(self, 'Успех', 'Данные были успешно изменены')
                else:
                    QMessageBox.critical(self, 'Ошибка', 'При изменения данных произошла ошибка', QMessageBox.Ok)
                    self.user_table()
            elif self.ui.tableWidget.columnCount() == 9:
                hardware_dict = {"Тип": 'type', "Название": 'name', "Описание": 'description'}

                if (self.ui.tableWidget.horizontalHeaderItem(current_col).text()) in hardware_dict:
                    col_name = hardware_dict[self.ui.tableWidget.horizontalHeaderItem(current_col).text()]
                    update_id = int(self.ui.tableWidget.item(current_row, 0).text())
                    hardware_response = helpers.patch_request('hardware', col_name, item, update_id)
                else:
                    col_name = 'specifications'
                    data = dict()
                    specs = ['log_elems', 'memory', 'pll', 'multiplier', 'pins']
                    for i in range(4, 9):
                        if self.ui.tableWidget.item(current_row, i) is not None:
                            if self.ui.tableWidget.item(current_row, i).text().isdigit():
                                data[specs[i - 4]] = int(self.ui.tableWidget.item(current_row, i).text())
                            else:
                                self.hardware_table()
                                return QMessageBox.critical(self, 'Ошибка',
                                                            'При изменения данных произошла ошибка.\nДопустим ввод только чисел.',
                                                            QMessageBox.Ok)

                    update_id = int(self.ui.tableWidget.item(current_row, 0).text())
                    hardware_response = helpers.patch_request('hardware', col_name, data, update_id)
                if hardware_response['detail'] == "OK":
                    QMessageBox.information(self, 'Успех', 'Данные были успешно изменены')
                else:
                    QMessageBox.critical(self, 'Ошибка', 'При изменения данных произошла ошибка', QMessageBox.Ok)
                    self.hardware_table()

    def request_table(self):
        self.ui.tableWidget.clearContents()
        self.ui.tableWidget.setColumnCount(12)
        self.ui.tableWidget.setHorizontalHeaderLabels(
            ["ID", "status", "location", "taken_date", "issued_by", "comment", "created", "user", "return_date",
             "hardware", "stock", "count"])

        reqs = helpers.get_request('request?joined=True')

        self.ui.tableWidget.setRowCount(len(reqs))
        for row in range(self.ui.tableWidget.rowCount()):
            self.ui.tableWidget.setRowHidden(row, False)
            self.status_box = QComboBox()
            statuses = ['new', 'taken', 'completed', 'canceled']
            self.status_box.addItems(statuses)
            self.boxes[row] = self.status_box
            current_status = reqs[row]["status"]
            index = statuses.index(current_status)
            self.status_box.setCurrentIndex(index)
            self.status_box.currentIndexChanged.connect(lambda _, row=row: self.patch_status(row))
            self.ui.tableWidget.setCellWidget(row, 1, self.status_box)
            self.ui.tableWidget.setItem(row, 2, QTableWidgetItem(reqs[row]["location"]))
            self.ui.tableWidget.setItem(row, 3, QTableWidgetItem(reqs[row]["taken_date"]))
            self.ui.tableWidget.setItem(row, 4, QTableWidgetItem(reqs[row]["issued_by"]))
            self.ui.tableWidget.setItem(row, 5, QTableWidgetItem(reqs[row]["comment"]))
            self.ui.tableWidget.setItem(row, 6, QTableWidgetItem(reqs[row]["created"]))
            self.ui.tableWidget.setItem(row, 7, QTableWidgetItem(reqs[row]["user"]))
            self.ui.tableWidget.setItem(row, 8, QTableWidgetItem(reqs[row]["return_date"]))
            self.ui.tableWidget.setItem(row, 0, QTableWidgetItem(str(reqs[row]["id"])))
            self.ui.tableWidget.setItem(row, 9, QTableWidgetItem(str(reqs[row]["hardware"])))
            self.ui.tableWidget.setItem(row, 10, QTableWidgetItem(str(reqs[row]["stock"])))
            self.ui.tableWidget.setItem(row, 11, QTableWidgetItem(str(reqs[row]["count"])))
            for col in range(2, 12):
                self.ui.tableWidget.item(row, col).setFlags(QtCore.Qt.ItemIsEnabled)

    def patch_status(self, tableRow):
        status = self.boxes[tableRow].currentText()
        print(status)
        col_name = 'status'
        update_id = int(self.ui.tableWidget.item(self.ui.tableWidget.currentRow(), 0).text())
        print(update_id)
        # request_response = patch_request('request', col_name, item, update_id)
        # if request_response['detail'] == "OK":
        #     QMessageBox.information(self, 'Успех', 'Данные были успешно изменены')
        # else:
        #     print(request_response)
        #     QMessageBox.critical(self, 'Ошибка', 'При изменения данных произошла ошибка', QMessageBox.Ok)

    def switch_theme(self):
        if self.styleSheet() != "":
            self.setStyleSheet("")
        else:
            self.setStyleSheet(qdarkstyle.load_stylesheet_pyqt5())

    def add_hardware(self):
        self.hard_dialog = HardDialog()
        self.hard_dialog.show()

    def add_user(self):
        self.user_dialog = UserDialog()
        self.user_dialog.show()

    def search_user(self):
        self.user_table()
        last_name = self.ui.search_line.text().lower()
        in_table = False
        for row in range(self.ui.tableWidget.rowCount()):
            if last_name == self.ui.tableWidget.item(row, 2).text().lower():
                in_table = True
            else:
                self.ui.tableWidget.setRowHidden(row, True)
        if not in_table:
            self.user_table()
            button = QMessageBox.question(
                self,
                'Ошибка',
                'Данного пользователя не существует. Хотите добавить?',
                QMessageBox.StandardButton.Yes |
                QMessageBox.StandardButton.No
            )
            if button == QMessageBox.StandardButton.Yes:
                self.add_user()

    def delete(self):
        current_row = self.ui.tableWidget.currentRow()
        if current_row < 0:
            return QMessageBox.warning(self, 'Внимание', 'Пожалуйста выберете запись для удаления')

        button = QMessageBox.question(
            self,
            'Подтверждение',
            'Вы уверены, что хотите удалить данную запись?',
            QMessageBox.StandardButton.Yes |
            QMessageBox.StandardButton.No
        )
        if button == QMessageBox.StandardButton.Yes:
            if self.ui.tableWidget.columnCount() == 8:
                delete_id = int(self.ui.tableWidget.item(current_row, 0).text())
                user_response = requests.delete(f"{DB_URL}/user/{delete_id}",
                                                headers={
                                                    'Authorization': DB_ACCESS_TOKEN}
                                                ).json()
                print(user_response)
                self.ui.tableWidget.removeRow(current_row)
            if self.ui.tableWidget.columnCount() == 9:
                delete_id = int(self.ui.tableWidget.item(current_row, 0).text())
                hardware_response = requests.delete(f"{DB_URL}/hardware/{delete_id}",
                                                    headers={
                                                        'Authorization': DB_ACCESS_TOKEN}
                                                    )
                hardware_user_response = hardware_response.json()
                print(hardware_user_response)
                self.ui.tableWidget.removeRow(current_row)

    def user_table(self):
        self.ui.tableWidget.clearContents()
        self.ui.tableWidget.setColumnCount(8)
        self.ui.tableWidget.setHorizontalHeaderLabels(
            ["ID", "Имя", "Фамилия", "Отчество", "Группа", "Уровень доступа", "Телефон", "Почта"])

        users = helpers.get_request('user')

        self.ui.tableWidget.setRowCount(len(users))
        for row in range(self.ui.tableWidget.rowCount()):
            self.ui.tableWidget.setRowHidden(row, False)
            self.ui.tableWidget.setItem(row, 1, QTableWidgetItem(users[row]["first_name"]))
            self.ui.tableWidget.setItem(row, 2, QTableWidgetItem(users[row]["last_name"]))
            self.ui.tableWidget.setItem(row, 3, QTableWidgetItem(users[row]["patronymic"]))
            self.ui.tableWidget.setItem(row, 4, QTableWidgetItem(users[row]["comment"]))
            self.ui.tableWidget.setItem(row, 5, QTableWidgetItem(users[row]["type"]))
            self.ui.tableWidget.setItem(row, 6, QTableWidgetItem(users[row]["phone"]))
            self.ui.tableWidget.setItem(row, 7, QTableWidgetItem(users[row]["email"]))
            self.ui.tableWidget.setItem(row, 0, QTableWidgetItem(str(users[row]["id"])))

    def hardware_table(self):
        self.ui.tableWidget.clearContents()
        self.ui.tableWidget.setColumnCount(9)
        self.ui.tableWidget.setHorizontalHeaderLabels(["ID", "Тип", "Название", "Описание", "log_elems",
                                                       "memory", "pll", "multiplier", "pins"
                                                       ])
        hardware = helpers.get_request('hardware')

        self.ui.tableWidget.setRowCount(len(hardware))
        self.ui.hardware_list.clear()
        for row in range(self.ui.tableWidget.rowCount()):
            self.ui.tableWidget.setRowHidden(row, False)
            self.ui.tableWidget.setItem(row, 1, QTableWidgetItem(hardware[row]["type"]))
            self.ui.tableWidget.setItem(row, 2, QTableWidgetItem(hardware[row]["name"]))
            self.ui.tableWidget.setItem(row, 3, QTableWidgetItem(hardware[row]["description"]))
            self.ui.tableWidget.setItem(row, 0, QTableWidgetItem(str(hardware[row]['id'])))

            specs = hardware[row]["specifications"]
            if specs:
                self.ui.tableWidget.setItem(row, 4, QTableWidgetItem(str(specs['log_elems'])))
                self.ui.tableWidget.setItem(row, 5, QTableWidgetItem(str(specs['memory'])))
                self.ui.tableWidget.setItem(row, 6, QTableWidgetItem(str(specs['pll'])))
                self.ui.tableWidget.setItem(row, 7, QTableWidgetItem(str(specs['multiplier'])))
                self.ui.tableWidget.setItem(row, 8, QTableWidgetItem(str(specs['pins'])))
            self.ui.hardware_list.addItem(QListWidgetItem(hardware[row]["name"]))

    def save_excel_file(self):
        """Создание эксель файла"""
        filepath, _ = QFileDialog.getSaveFileName(self, 'Save File', '', ".xls(*.xls)")
        wbk = Workbook()
        sheet = wbk.add_sheet("sheet", cell_overwrite_ok=True)
        self.make_sheet(sheet)
        wbk.save(filepath)

    def make_sheet(self, sheet):
        """Создание Листа в экселе"""
        for currentColumn in range(self.ui.tableWidget.columnCount()):
            sheet.write(0, currentColumn, str(self.ui.tableWidget.horizontalHeaderItem(currentColumn).text()))
        for currentColumn in range(self.ui.tableWidget.columnCount()):
            for currentRow in range(self.ui.tableWidget.rowCount()):
                try:
                    teext = str(self.ui.tableWidget.item(currentRow, currentColumn).text())
                    sheet.write(currentRow + 1, currentColumn, teext)
                except AttributeError:
                    pass

    def save_doc_file(self):
        """Создание док файла"""
        # получаем путь к файлу
        filepath, _ = QFileDialog.getSaveFileName(self, 'Save File', '', ".docx(*.docx)")
        # добавляем таблицу
        doc = Document()
        table = doc.add_table(rows=self.ui.tableWidget.rowCount() + 1, cols=self.ui.tableWidget.columnCount())
        # применяем стиль для таблицы
        table.style = 'Table Grid'
        # добавляем заголовки
        for currentColumn in range(self.ui.tableWidget.columnCount()):
            cell = table.cell(0, currentColumn)
            cell.text = str(self.ui.tableWidget.horizontalHeaderItem(currentColumn).text())
        # заполняем таблицу данными
        for row in range(self.ui.tableWidget.rowCount()):
            for col in range(self.ui.tableWidget.columnCount()):
                cell = table.cell(row + 1, col)
                # записываем в ячейку данные
                try:
                    cell.text = str(self.ui.tableWidget.item(row, col).text())
                except AttributeError:
                    pass
        doc.save(filepath)


if __name__ == '__main__':
    app = QtWidgets.QApplication(sys.argv)
    window = LoginWindow()
    window.show()
    sys.exit(app.exec_())

"""
    error = QMessageBox()
    error.setWindowTitle("Ошибка")
    error.setText("Введён неверный пароль")
    error.setIcon(QMessageBox.Icon.Ok)
    error.setStandardButtons(QMessageBox.Ok)

    error.exec_()
"""
