import json
import sys

import qdarkstyle
from PyQt5 import QtWidgets
from PyQt5.QtWidgets import QFileDialog, QMessageBox, QTableWidgetItem, QListWidgetItem, QDialog, QDialogButtonBox
from entry_ui import Ui_Enter
from window_ui import Ui_MainWindow
from user_ui import Ui_Newuser
from hardware_ui import Ui_Newhard
from xlwt import Workbook
from docx import Document
import requests

CODE = '1'
DB_ACCESS_TOKEN = "Basic NVJOWUJkTGR1VER4UUNjTThZWXJiNW5BOkg0ZFNjQXlHYlM4OUtnTGdaQnMydlBzaw=="
DB_URL = "https://helow19274.ru/aip/api"


class LoginWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.table_window = Ui_MainWindow
        self.ui = Ui_Enter()
        self.ui.setupUi(self)
        self.ui.btn_entry.clicked.connect(self.on_click)

    def on_click(self):
        """Проверка пароля"""
        password = self.ui.line_password.text()
        if len(password) > 0:
            if password == CODE:
                self.close()
                self.table_window = MainWindow()
                self.table_window.show()
            else:
                self.ui.label.setText('Wrong password!')
                self.ui.line_password.clear()


class HardDialog(QDialog):
    def __init__(self):
        super().__init__()
        self.ui = Ui_Newhard()
        self.ui.setupUi(self)
        self.setStyleSheet(open('hard_style.qss').read())
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.exit)

    def exit(self):
        self.ui.name_line.setText('')
        self.close()

    def check(self):
        name = self.ui.name_line.text()
        specs = self.ui.spec_line.text()
        count = self.ui.count_line.text()
        description = self.ui.description_line.text()
        typ = self.ui.type_line.text()
        if typ == "":
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе типа платы",
                QMessageBox.StandardButton.Ok
            )
        elif name == "":
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе названия",
                QMessageBox.StandardButton.Ok
            )
        elif description == "":
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе описания. Если оно отсутствует введите -",
                QMessageBox.StandardButton.Ok
            )
        elif specs == "":
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе спецификации",
                QMessageBox.StandardButton.Ok
            )
        elif count == "" or not count.isdigit():
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе количества",
                QMessageBox.StandardButton.Ok
            )
        else:
            self.close()


class UserDialog(QDialog):
    def __init__(self):
        super().__init__()
        self.ui = Ui_Newuser()
        self.ui.setupUi(self)
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.exit)

    def exit(self):
        self.ui.surname_line.setText('')
        self.close()

    def check(self):
        surname = self.ui.surname_line.text()
        name = self.ui.name_line.text()
        patr = self.ui.patr_line.text()
        group = self.ui.group_line.text()
        email = self.ui.email_line.text()
        phone = self.ui.phone_line.text()
        if surname == "":
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе фамилии",
                QMessageBox.StandardButton.Ok
            )
        elif name == "":
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе имени",
                QMessageBox.StandardButton.Ok
            )
        elif patr == "":
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе отчества. Если оно отсутствует введите -",
                QMessageBox.StandardButton.Ok
            )
        elif group == "":
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе группы",
                QMessageBox.StandardButton.Ok
            )
        elif email == "" or "@" not in email:
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе почты",
                QMessageBox.StandardButton.Ok
            )
        elif phone == "" or len(phone) != 11:
            QMessageBox.warning(
                self,
                "Ошибка",
                "Ошибка в вводе телефона",
                QMessageBox.StandardButton.Ok
            )
        else:
            self.close()


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.hard_dialog = Ui_Newhard
        self.user_dialog = Ui_Newuser
        self.ui = Ui_MainWindow()
        self.ui.setupUi(self)
        self.ui.action_excel.triggered.connect(self.save_excel_file)
        self.ui.action_doc.triggered.connect(self.save_doc_file)

        self.ui.users_button.clicked.connect(self.user_table)
        self.ui.hardware_button.clicked.connect(self.hardware_table)
        self.ui.request_button.clicked.connect(self.request_table)

        self.ui.add_user_button.clicked.connect(self.add_user)
        self.ui.add_hardware_button.clicked.connect(self.add_hard)

        self.ui.search_button.clicked.connect(self.search_user)
        self.ui.del_button.clicked.connect(self.delete)

        self.ui.action_theme.triggered.connect(self.switch_theme)

    def request_table(self):
        self.ui.tableWidget.clearContents()
        self.ui.tableWidget.setColumnCount(12)
        self.ui.tableWidget.setHorizontalHeaderLabels(
            ["status", "location", "taken_date", "issued_by", "comment", "created", "user", "return_date", "id",
             "hardware", "stock", "count"])

        reqs = requests.get(f"{DB_URL}/request",
                            headers={
                                'Authorization': DB_ACCESS_TOKEN}).json()

        self.ui.tableWidget.setRowCount(len(reqs))
        for row in range(self.ui.tableWidget.rowCount()):
            self.ui.tableWidget.setRowHidden(row, False)
            self.ui.tableWidget.setItem(row, 0, QTableWidgetItem(reqs[row]["status"]))
            self.ui.tableWidget.setItem(row, 1, QTableWidgetItem(reqs[row]["location"]))
            self.ui.tableWidget.setItem(row, 2, QTableWidgetItem(reqs[row]["taken_date"]))
            self.ui.tableWidget.setItem(row, 3, QTableWidgetItem(reqs[row]["issued_by"]))
            self.ui.tableWidget.setItem(row, 4, QTableWidgetItem(reqs[row]["comment"]))
            self.ui.tableWidget.setItem(row, 5, QTableWidgetItem(reqs[row]["created"]))
            self.ui.tableWidget.setItem(row, 6, QTableWidgetItem(reqs[row]["user"]))
            self.ui.tableWidget.setItem(row, 7, QTableWidgetItem(reqs[row]["return_date"]))
            self.ui.tableWidget.setItem(row, 8, QTableWidgetItem(str(reqs[row]["id"])))
            self.ui.tableWidget.setItem(row, 9, QTableWidgetItem(str(reqs[row]["hardware"])))
            self.ui.tableWidget.setItem(row, 10, QTableWidgetItem(str(reqs[row]["stock"])))
            self.ui.tableWidget.setItem(row, 11, QTableWidgetItem(str(reqs[row]["count"])))

    def switch_theme(self):
        if self.styleSheet() != "":
            self.setStyleSheet("")
        else:
            self.setStyleSheet(qdarkstyle.load_stylesheet_pyqt5())

    def add_hard(self):
        self.hard_dialog = HardDialog()
        self.hard_dialog.show()
        self.hard_dialog.exec_()
        name = self.hard_dialog.ui.name_line.text()
        if name != '':
            specs = self.hard_dialog.ui.spec_line.text()
            count = self.hard_dialog.ui.count_line.text()
            description = self.hard_dialog.ui.description_line.text()
            typ = self.hard_dialog.ui.type_line.text()
            hardware_data = {
                "name": name,
                "type": typ,
                "description": description,
                "image_link": "https://encrypted-tbn0.gstatic.com/images?q=tbn:ANd9GcStW9UvyhB2beq-tiyJMhzWdP98Rny7PzRaPA&usqp=CAU",
                "specifications": specs
            }
            hardware_request_body = json.dumps(hardware_data, ensure_ascii=False)

            hardware_response = requests.post(f"{DB_URL}/hardware",
                                              headers={
                                                  'Authorization': DB_ACCESS_TOKEN,
                                                  'accept': 'application/json',
                                                  'Content-Type': 'application/json'
                                              },
                                              data=hardware_request_body
                                              )
            hardware_user_response = hardware_response.json()
            print(hardware_user_response)
        else:
            print("Была отмена")

    def add_user(self):
        self.user_dialog = UserDialog()
        self.user_dialog.show()
        self.user_dialog.exec_()
        surname = self.user_dialog.ui.surname_line.text()
        if surname != '':
            name = self.user_dialog.ui.name_line.text()
            patr = self.user_dialog.ui.patr_line.text()
            group = self.user_dialog.ui.group_line.text()
            email = self.user_dialog.ui.email_line.text()
            phone = self.user_dialog.ui.phone_line.text()
            access = self.user_dialog.ui.access_box.currentText()
            user_data = {
                "active": True,
                "type": access,
                "first_name": name,
                "last_name": surname,
                "patronymic": patr,
                "image_link": "https://clck.ru/34TfSF",
                "email": email,
                "phone": phone,
                "card_id": "string",
                "card_key": "string",
                "comment": group
            }

            user_request_body = json.dumps(user_data, ensure_ascii=False)

            response = requests.post(f"{DB_URL}/user",
                                     headers={'Authorization': DB_ACCESS_TOKEN},
                                     data=user_request_body
                                     )
            user_response = response.json()
            print(user_response)
        else:
            print("Была отмена")

    def search_user(self):
        self.user_table()
        last_name = self.ui.search_line.text().lower()
        in_table = False
        for row in range(self.ui.tableWidget.rowCount()):
            if last_name == self.ui.tableWidget.item(row, 1).text().lower():
                in_table = True
            else:
                self.ui.tableWidget.setRowHidden(row, True)
        if not in_table:
            self.user_table()
            button = QMessageBox.question(
                self,
                'Ошибка',
                'Данного пользователя не существует. Хотите добавить?',
                QMessageBox.StandardButton.Yes |
                QMessageBox.StandardButton.No
            )
            if button == QMessageBox.StandardButton.Yes:
                self.add_user()

    def delete(self):
        current_row = self.ui.tableWidget.currentRow()
        if current_row < 0:
            return QMessageBox.warning(self, 'Внимание', 'Пожалуйста выберете запись для удаления')

        button = QMessageBox.question(
            self,
            'Подтверждение',
            'Вы уверены, что хотите удалить данную запись?',
            QMessageBox.StandardButton.Yes |
            QMessageBox.StandardButton.No
        )
        if button == QMessageBox.StandardButton.Yes:
            if self.ui.tableWidget.columnCount() == 8:
                delete_id = int(self.ui.tableWidget.item(current_row, 7).text())
                user_response = requests.delete(f"{DB_URL}/user/{delete_id}",
                                                headers={
                                                    'Authorization': DB_ACCESS_TOKEN}
                                                ).json()
                print(user_response)
                self.ui.tableWidget.removeRow(current_row)
            if self.ui.tableWidget.columnCount() == 5:
                delete_id = int(self.ui.tableWidget.item(current_row, 4).text())
                hardware_response = requests.delete(f"{DB_URL}/hardware/{delete_id}",
                                                    headers={
                                                        'Authorization': DB_ACCESS_TOKEN}
                                                    )
                hardware_user_response = hardware_response.json()
                print(hardware_user_response)
                self.ui.tableWidget.removeRow(current_row)

    def user_table(self):
        self.ui.tableWidget.clearContents()
        self.ui.tableWidget.setColumnCount(8)
        self.ui.tableWidget.setHorizontalHeaderLabels(
            ["Имя", "Фамилия", "Отчество", "Группа", "Уровень доступа", "Телефон", "Почта", "ID"])

        users = requests.get(f"{DB_URL}/user",
                             headers={
                                 'Authorization': DB_ACCESS_TOKEN}).json()

        self.ui.tableWidget.setRowCount(len(users))
        for row in range(self.ui.tableWidget.rowCount()):
            self.ui.tableWidget.setRowHidden(row, False)
            self.ui.tableWidget.setItem(row, 0, QTableWidgetItem(users[row]["first_name"]))
            self.ui.tableWidget.setItem(row, 1, QTableWidgetItem(users[row]["last_name"]))
            self.ui.tableWidget.setItem(row, 2, QTableWidgetItem(users[row]["patronymic"]))
            self.ui.tableWidget.setItem(row, 3, QTableWidgetItem(users[row]["comment"]))
            self.ui.tableWidget.setItem(row, 4, QTableWidgetItem(users[row]["type"]))
            self.ui.tableWidget.setItem(row, 5, QTableWidgetItem(users[row]["phone"]))
            self.ui.tableWidget.setItem(row, 6, QTableWidgetItem(users[row]["email"]))
            self.ui.tableWidget.setItem(row, 7, QTableWidgetItem(str(users[row]["id"])))

    def add_row(self):
        row_index_to_insert = self.ui.tableWidget.rowCount() + 1
        self.ui.tableWidget.setRowCount(row_index_to_insert)

        for column in range(self.ui.tableWidget.columnCount()):
            item = QTableWidgetItem('')
            # item.setFlags(item.flags() & ~Qt.ItemIsEditable)
            self.ui.tableWidget.setItem(row_index_to_insert - 1, column, item)

    def hardware_table(self):
        self.ui.tableWidget.clearContents()
        self.ui.tableWidget.setColumnCount(5)
        self.ui.tableWidget.setHorizontalHeaderLabels(["Тип", "Название", "Описание", "Спецификация", "ID"])

        hardware = requests.get(f"{DB_URL}/hardware",
                                headers={
                                    'Authorization': DB_ACCESS_TOKEN}).json()

        self.ui.tableWidget.setRowCount(len(hardware))
        self.ui.hardware_list.clear()
        for row in range(self.ui.tableWidget.rowCount()):
            self.ui.tableWidget.setRowHidden(row, False)
            self.ui.tableWidget.setItem(row, 0, QTableWidgetItem(hardware[row]["type"]))
            self.ui.tableWidget.setItem(row, 1, QTableWidgetItem(hardware[row]["name"]))
            self.ui.tableWidget.setItem(row, 2, QTableWidgetItem(hardware[row]["description"]))
            self.ui.tableWidget.setItem(row, 4, QTableWidgetItem(str(hardware[row]['id'])))

            specs = str(hardware[row]["specifications"]).replace('{','').replace('}','').replace(',', '\n')
            self.ui.tableWidget.setItem(row, 3, QTableWidgetItem(specs))
            self.ui.hardware_list.addItem(QListWidgetItem(hardware[row]["name"]))

    def save_excel_file(self):
        """Создание эксель файла"""
        filepath, _ = QFileDialog.getSaveFileName(self, 'Save File', '', ".xls(*.xls)")
        wbk = Workbook()
        sheet = wbk.add_sheet("sheet", cell_overwrite_ok=True)
        self.make_sheet(sheet)
        wbk.save(filepath)

    def make_sheet(self, sheet):
        """Создание Листа в экселе"""
        for currentColumn in range(self.ui.tableWidget.columnCount()):
            sheet.write(0, currentColumn, str(self.ui.tableWidget.horizontalHeaderItem(currentColumn).text()))
        for currentColumn in range(self.ui.tableWidget.columnCount()):
            for currentRow in range(self.ui.tableWidget.rowCount()):
                try:
                    teext = str(self.ui.tableWidget.item(currentRow, currentColumn).text())
                    sheet.write(currentRow + 1, currentColumn, teext)
                except AttributeError:
                    pass

    def save_doc_file(self):
        """Создание док файла"""
        # получаем путь к файлу
        filepath, _ = QFileDialog.getSaveFileName(self, 'Save File', '', ".docx(*.docx)")
        # добавляем таблицу
        doc = Document()
        table = doc.add_table(rows=self.ui.tableWidget.rowCount() + 1, cols=self.ui.tableWidget.columnCount())
        # применяем стиль для таблицы
        table.style = 'Table Grid'
        # добавляем заголовки
        for currentColumn in range(self.ui.tableWidget.columnCount()):
            cell = table.cell(0, currentColumn)
            cell.text = str(self.ui.tableWidget.horizontalHeaderItem(currentColumn).text())
        # заполняем таблицу данными
        for row in range(self.ui.tableWidget.rowCount()):
            for col in range(self.ui.tableWidget.columnCount()):
                cell = table.cell(row + 1, col)
                # записываем в ячейку данные
                try:
                    cell.text = str(self.ui.tableWidget.item(row, col).text())
                except AttributeError:
                    pass
        doc.save(filepath)


if __name__ == '__main__':
    app = QtWidgets.QApplication(sys.argv)
    window = LoginWindow()
    window.show()
    sys.exit(app.exec_())

''' 
    error = QMessageBox()
    error.setWindowTitle("Ошибка")
    error.setText("Введён неверный пароль")
    error.setIcon(QMessageBox.Warning)
    error.setStandardButtons(QMessageBox.Ok)

    error.exec_()
'''
